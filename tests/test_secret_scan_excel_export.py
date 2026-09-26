from __future__ import annotations

import os
import sys

import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from scanindex.ui.screens.secret_file_scan_screen import (
    EXCEL_HEADERS,
    SecretScanMatch,
    _apply_kie_org_fallback,
    _declass_cell_text,
    _file_mtime_for_display,
    _format_file_mtime,
    _kie_org_fallback_name,
    _org_code_from_filename,
    _org_display_name,
    _org_name_lookup,
    export_matches_to_excel,
    load_secret_matches_from_excel,
)
from scanindex.infra.version import get_version


def _match(
    keyword: str = "Mật",
    page: int = 1,
    issue_year: int = 0,
    declass_years: int = 0,
    declass_due: bool = False,
    source_version: str = "",
    source_path: str = r"D:\tailieu\hocbao\congvan_01.pdf",
    kie_org_name: str = "",
    file_mtime: float = 0.0,
) -> SecretScanMatch:
    return SecretScanMatch(
        source_path=source_path,
        relative_path=os.path.join("hocbao", os.path.basename(source_path)),
        keyword=keyword,
        page_number=page,
        mode="Trang đầu",
        ocr_pdf_path="",
        note="PDF gốc; trang đầu",
        issue_year=issue_year,
        declass_years=declass_years,
        declass_due=declass_due,
        source_version=source_version,
        kie_org_name=kie_org_name,
        file_mtime=file_mtime,
    )


def test_export_matches_to_excel_writes_all_fields(tmp_path) -> None:
    import openpyxl

    dest = str(tmp_path / "danh_sach.xlsx")
    matches = [_match("Mật", 1), _match("Tuyệt mật", 2)]
    export_matches_to_excel(matches, dest)

    wb = openpyxl.load_workbook(dest)
    ws = wb.active
    assert ws.title == "Van ban mat"
    headers = [ws.cell(row=1, column=c).value for c in range(1, len(EXCEL_HEADERS) + 1)]
    assert headers == EXCEL_HEADERS
    assert EXCEL_HEADERS[3] == "Mã cơ quan"
    assert EXCEL_HEADERS[4] == "Tên cơ quan"
    assert EXCEL_HEADERS[7] == "Ngày cập nhật"
    assert EXCEL_HEADERS[10] == "Đáp ứng giải mật"
    assert EXCEL_HEADERS[11] == "Ghi chú"

    assert ws.cell(row=2, column=1).value == 1
    assert ws.cell(row=2, column=2).value == "Mật"
    assert ws.cell(row=2, column=3).value == "congvan_01.pdf"
    # Tên file không dạng <uuid>_<mã>-… → 2 cột mã/tên cơ quan trống.
    assert ws.cell(row=2, column=4).value in ("", None)
    assert ws.cell(row=2, column=5).value in ("", None)
    assert os.path.basename(ws.cell(row=2, column=6).value) == "congvan_01.pdf"
    assert ws.cell(row=2, column=7).value == r"D:\tailieu\hocbao\congvan_01.pdf"
    # File không tồn tại trên đĩa và không có mtime dự phòng → ô trống.
    assert ws.cell(row=2, column=8).value in ("", None)
    assert ws.cell(row=2, column=9).value == 1
    assert ws.cell(row=2, column=10).value == "Trang đầu"
    # Chưa đạt giải mật → ô trống.
    assert ws.cell(row=2, column=11).value in ("", None)
    assert ws.cell(row=2, column=12).value == "PDF gốc; trang đầu"

    assert ws.cell(row=3, column=1).value == 2
    assert ws.cell(row=3, column=2).value == "Tuyệt mật"
    assert ws.cell(row=3, column=9).value == 2


def test_declass_cell_text_two_states_only() -> None:
    current = get_version()
    due = _match(issue_year=2016, declass_years=10, declass_due=True,
                 source_version=current)
    assert _declass_cell_text(due, current) == "Đáp ứng (2016, 10 năm)"
    # Mọi trường hợp khác → trống (chưa đạt).
    not_due = _match(issue_year=2025, declass_years=20, source_version=current)
    assert _declass_cell_text(not_due, current) == ""
    inherited = _match()  # không năm, không source_version
    assert _declass_cell_text(inherited, current) == ""
    fresh_no_year = _match(source_version=current)
    assert _declass_cell_text(fresh_no_year, current) == ""


# Ví dụ của người dùng: mã P9 (không có tên) và A29.183 (Đảng ủy Xã Hưng Long).
_UUID_A29 = (
    r"D:\kho\39cfea63-af46-460f-8c1c-340ab5e3c6af_"
    r"A29.183-A29.37.28.001-04-0010-082.pdf"
)
_UUID_P9 = (
    r"D:\kho\4bded843-8b83-4d68-a3de-32973498177e_P9-16-04-10.pdf"
)


def test_org_code_from_filename() -> None:
    assert _org_code_from_filename(_UUID_A29) == "A29.183"
    assert _org_code_from_filename(_UUID_P9) == "P9"
    # Không dạng uuid → trống.
    assert _org_code_from_filename(r"D:\kho\congvan_01.pdf") == ""
    assert _org_code_from_filename("") == ""


def test_org_name_lookup_examples() -> None:
    lookup = _org_name_lookup()
    assert lookup.get("A29.183") == "Đảng ủy Xã Hưng Long"
    # P9 không có trong bảng → tra không ra (cột Tên cơ quan trống).
    assert "P9" not in lookup


def test_org_lookup_finds_bundled_and_prefers_exe_dir(tmp_path, monkeypatch) -> None:
    """Bản đóng gói onedir đặt data trong _internal — code phải tìm được qua
    get_resource_path; file đặt cạnh exe được ưu tiên để thay danh mục."""
    import json as jsonlib

    from scanindex.ui.screens import secret_file_scan_screen as sfss

    bundle_dir = tmp_path / "_internal"
    bundle_dir.mkdir()
    (bundle_dir / "madinhdanh_lookup.json").write_text(
        jsonlib.dumps([{"ma_dinh_danh": "T99", "ten_co_quan": "Bản trong bundle"}]),
        encoding="utf-8",
    )
    exe_dir = tmp_path / "app"
    exe_dir.mkdir()

    def fake_resource_path(rel: str) -> str:
        # Giống get_resource_path thật: cạnh exe trước, bundle sau.
        beside = exe_dir / rel
        if beside.exists():
            return str(beside)
        return str(bundle_dir / rel)

    monkeypatch.setattr(sfss, "get_resource_path", fake_resource_path)
    monkeypatch.setattr(sfss, "_ORG_LOOKUP_CACHE", None)
    assert sfss._org_name_lookup().get("T99") == "Bản trong bundle"

    # Đặt file mới cạnh exe → override bản đóng gói (không cần build lại).
    (exe_dir / "madinhdanh_lookup.json").write_text(
        jsonlib.dumps([{"ma_dinh_danh": "T99", "ten_co_quan": "Bản cạnh exe"}]),
        encoding="utf-8",
    )
    monkeypatch.setattr(sfss, "_ORG_LOOKUP_CACHE", None)
    assert sfss._org_name_lookup().get("T99") == "Bản cạnh exe"
    monkeypatch.setattr(sfss, "_ORG_LOOKUP_CACHE", None)


def test_export_writes_org_columns(tmp_path) -> None:
    import openpyxl

    dest = str(tmp_path / "danh_sach.xlsx")
    m1 = _match(source_path=_UUID_A29)
    m2 = _match(keyword="Tuyệt mật", source_path=_UUID_P9)
    export_matches_to_excel([m1, m2], dest)

    ws = openpyxl.load_workbook(dest).active
    assert ws.cell(row=2, column=4).value == "A29.183"
    assert ws.cell(row=2, column=5).value == "Đảng ủy Xã Hưng Long"
    assert ws.cell(row=3, column=4).value == "P9"
    assert ws.cell(row=3, column=5).value in ("", None)


def test_export_creates_parent_dirs(tmp_path) -> None:
    dest = str(tmp_path / "sub" / "dir" / "out.xlsx")
    export_matches_to_excel([_match()], dest)
    assert os.path.isfile(dest)


def test_export_empty_list_still_writes_header(tmp_path) -> None:
    import openpyxl

    dest = str(tmp_path / "empty.xlsx")
    export_matches_to_excel([], dest)
    ws = openpyxl.load_workbook(dest).active
    assert ws.max_row == 1
    assert ws.cell(row=1, column=1).value == "STT"


# ---------------------------------------------------------------------------
# Tên cơ quan dự phòng KIE: "<cơ quan ban hành> <cơ quan cấp trên>*"
# ---------------------------------------------------------------------------

def _field(label: str, text: str, page_index: int = 0) -> dict:
    return {"label": label, "page_index": page_index, "text": text}


def test_kie_org_fallback_name_both_orgs(tmp_path, monkeypatch) -> None:
    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        lambda path, page: [
            _field("ISSUE_ORG_SUPERIOR", "ỦY BAN NHÂN DÂN TỈNH BẮC NINH"),
            _field("ISSUE_ORG_NAME", "SỞ TƯ PHÁP"),
        ],
    )
    # Ban hành trước, cấp trên sau, "*" ở cuối.
    assert (
        _kie_org_fallback_name(str(canonical), 0)
        == "SỞ TƯ PHÁP ỦY BAN NHÂN DÂN TỈNH BẮC NINH*"
    )


def test_kie_org_fallback_name_partial_and_multiline(tmp_path, monkeypatch) -> None:
    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        lambda path, page: [
            # Field multi_line nối dòng bằng "\n" — phải gộp về 1 dòng.
            _field("ISSUE_ORG_NAME", "ỦY BAN NHÂN DÂN\n     XÃ HƯNG LONG"),
            _field("PLACE_DATE", "Bắc Ninh, ngày 04 tháng 4 năm 2016"),
        ],
    )
    assert _kie_org_fallback_name(str(canonical), 0) == "ỦY BAN NHÂN DÂN XÃ HƯNG LONG*"


def test_kie_org_fallback_name_no_orgs(tmp_path, monkeypatch) -> None:
    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        lambda path, page: [_field("PLACE_DATE", "ngày 04 tháng 4 năm 2016")],
    )
    assert _kie_org_fallback_name(str(canonical), 0) == ""
    # Model lỗi / không có sẵn (None) → trống, không raise.
    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        lambda path, page: None,
    )
    assert _kie_org_fallback_name(str(canonical), 0) == ""


def test_apply_kie_org_fallback_only_when_lookup_misses(tmp_path, monkeypatch) -> None:
    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    calls: list[int] = []

    def fake_fields(path: str, page: int) -> list[dict]:
        calls.append(page)
        # Field phải nằm đúng trang chứa dấu mới được tính (page 3 → index 2).
        return [
            _field(
                "ISSUE_ORG_NAME",
                "SỞ TÀI NGUYÊN VÀ MÔI TRƯỜNG",
                page_index=page,
            )
        ]

    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        fake_fields,
    )
    resolved = _match(source_path=_UUID_A29)  # A29.183 tra được tên
    unresolved = _match(keyword="Tuyệt mật", page=3, source_path=_UUID_P9)
    _apply_kie_org_fallback([resolved, unresolved], str(canonical))

    # Tra được tên → không đụng KIE, không ghi đè.
    assert resolved.kie_org_name == ""
    # P9 không có trong danh mục → điền dự phòng, theo trang của dấu (3-1=2).
    assert unresolved.kie_org_name == "SỞ TÀI NGUYÊN VÀ MÔI TRƯỜNG*"
    assert calls == [2]


def test_apply_kie_org_fallback_env_switch(tmp_path, monkeypatch) -> None:
    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    monkeypatch.setenv("SECRET_SCAN_DISABLE_KIE_ORG", "1")
    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        lambda path, page: pytest.fail("phải tắt hẳn, không gọi KIE"),
    )
    m = _match(source_path=_UUID_P9)
    _apply_kie_org_fallback([m], str(canonical))
    assert m.kie_org_name == ""


def test_apply_kie_org_fallback_missing_canonical(tmp_path) -> None:
    m = _match(source_path=_UUID_P9)
    _apply_kie_org_fallback([m], str(tmp_path / "khong_co.json"))
    assert m.kie_org_name == ""


def test_kie_fields_cache_dedupes_inference(tmp_path, monkeypatch) -> None:
    """Pass năm và pass tên cơ quan cùng trang → chỉ 1 lần infer."""
    from scanindex.core.kie import engine as kie_engine
    from scanindex.ui.screens import secret_file_scan_screen as sfss

    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    calls = []

    def fake_run(canonical_json_path, selected_pages=None):
        calls.append(selected_pages)
        return {
            "field_instances": [
                _field("PLACE_DATE", "ngày 04 tháng 4 năm 2016"),
                _field("ISSUE_ORG_NAME", "SỞ TƯ PHÁP"),
            ]
        }

    monkeypatch.setattr(kie_engine, "_run_layoutlmv3", fake_run)
    sfss._KIE_FIELDS_CACHE.clear()
    try:
        fields1 = sfss._kie_fields_for_page(str(canonical), 0)
        fields2 = sfss._kie_fields_for_page(str(canonical), 0)
        assert fields1 is fields2
        assert len(fields1) == 2
        assert calls == [[0]]  # lần 2 ăn cache
    finally:
        sfss._KIE_FIELDS_CACHE.clear()


def test_org_display_name_prefers_lookup() -> None:
    # Tra được qua mã → dùng tên danh mục, bỏ dự phòng KIE.
    assert (
        _org_display_name(_UUID_A29, "SỞ GIAO THÔNG*")
        == "Đảng ủy Xã Hưng Long"
    )
    # Không tra được → dùng dự phòng KIE (có "*").
    assert (
        _org_display_name(_UUID_P9, "SỞ GIAO THÔNG VẬN TẢI*")
        == "SỞ GIAO THÔNG VẬN TẢI*"
    )
    # Không mã, không dự phòng → trống.
    assert _org_display_name(r"D:\kho\congvan_01.pdf", "") == ""


def test_export_writes_kie_org_fallback(tmp_path) -> None:
    import openpyxl

    dest = str(tmp_path / "danh_sach.xlsx")
    m1 = _match(source_path=_UUID_P9, kie_org_name="SỞ CÔNG THƯƠNG BỘ CÔNG THƯƠNG*")
    # Lookup vẫn thắng khi mã tra được, kể cả khi có dự phòng.
    m2 = _match(
        keyword="Tuyệt mật",
        source_path=_UUID_A29,
        kie_org_name="SỞ CÔNG THƯƠNG BỘ CÔNG THƯƠNG*",
    )
    export_matches_to_excel([m1, m2], dest)

    ws = openpyxl.load_workbook(dest).active
    assert ws.cell(row=2, column=4).value == "P9"
    assert ws.cell(row=2, column=5).value == "SỞ CÔNG THƯƠNG BỘ CÔNG THƯƠNG*"
    assert ws.cell(row=3, column=4).value == "A29.183"
    assert ws.cell(row=3, column=5).value == "Đảng ủy Xã Hưng Long"


def test_load_excel_round_trips_kie_org_fallback(tmp_path) -> None:
    dest = str(tmp_path / "danh_sach.xlsx")
    m = _match(source_path=_UUID_P9, kie_org_name="SỞ CÔNG THƯƠNG BỘ CÔNG THƯƠNG*")
    export_matches_to_excel([m], dest)

    loaded = load_secret_matches_from_excel(dest)
    assert len(loaded) == 1
    assert loaded[0].kie_org_name == "SỞ CÔNG THƯƠNG BỘ CÔNG THƯƠNG*"
    # Hiển thị sau khi nạp lại vẫn ra đúng tên dự phòng.
    assert (
        _org_display_name(loaded[0].source_path, loaded[0].kie_org_name)
        == "SỞ CÔNG THƯƠNG BỘ CÔNG THƯƠNG*"
    )


def test_load_excel_does_not_keep_catalog_names_as_fallback(tmp_path) -> None:
    """Tên tra từ danh mục (không "*") không được giữ làm dự phòng khi nạp
    lại — lỡ danh mục đổi thì tra lại, không treo bản sao cũ."""
    dest = str(tmp_path / "danh_sach.xlsx")
    m = _match(source_path=_UUID_A29)  # cột Tên cơ quan = tên danh mục
    export_matches_to_excel([m], dest)

    loaded = load_secret_matches_from_excel(dest)
    assert loaded[0].kie_org_name == ""


# ---------------------------------------------------------------------------
# Exception: KIE lỗi / payload rác không được phép kéo sập lượt quét
# ---------------------------------------------------------------------------

def test_kie_fields_survives_garbage_payload(tmp_path, monkeypatch) -> None:
    from scanindex.core.kie import engine as kie_engine
    from scanindex.ui.screens import secret_file_scan_screen as sfss

    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")

    # Model trả payload không phải dict → None, không raise.
    monkeypatch.setattr(
        kie_engine, "_run_layoutlmv3", lambda p, selected_pages=None: ["junk"]
    )
    sfss._KIE_FIELDS_CACHE.clear()
    try:
        assert sfss._kie_fields_for_page(str(canonical), 0) is None
        # Infer lỗi được cache âm: trang này không đập model liên tục.
        assert sfss._kie_cached_fields(str(canonical), 0) is None
    finally:
        sfss._KIE_FIELDS_CACHE.clear()

    # field_instances lẫn rác (chuỗi/None/số) → chỉ giữ dict hợp lệ.
    monkeypatch.setattr(
        kie_engine,
        "_run_layoutlmv3",
        lambda p, selected_pages=None: {
            "field_instances": [
                "x",
                None,
                42,
                {"label": "ISSUE_ORG_NAME", "page_index": 0, "text": "SỞ Y TẾ"},
            ]
        },
    )
    sfss._KIE_FIELDS_CACHE.clear()
    try:
        assert sfss._kie_fields_for_page(str(canonical), 0) == [
            {"label": "ISSUE_ORG_NAME", "page_index": 0, "text": "SỞ Y TẾ"}
        ]
    finally:
        sfss._KIE_FIELDS_CACHE.clear()

    # Model raise (ORT chết, thiếu model…) → None, không ném ra ngoài.
    def boom(p, selected_pages=None):
        raise RuntimeError("ORT dead")

    monkeypatch.setattr(kie_engine, "_run_layoutlmv3", boom)
    sfss._KIE_FIELDS_CACHE.clear()
    try:
        assert sfss._kie_fields_for_page(str(canonical), 0) is None
    finally:
        sfss._KIE_FIELDS_CACHE.clear()


def test_kie_cache_miss_sentinel(tmp_path, monkeypatch) -> None:
    from scanindex.core.kie import engine as kie_engine
    from scanindex.ui.screens import secret_file_scan_screen as sfss

    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(
        kie_engine,
        "_run_layoutlmv3",
        lambda p, selected_pages=None: {"field_instances": []},
    )
    sfss._KIE_FIELDS_CACHE.clear()
    try:
        # Chưa infer → sentinel MISS (khác None của "infer rồi nhưng lỗi").
        assert (
            sfss._kie_cached_fields(str(canonical), 0) is sfss._KIE_CACHE_MISS
        )
        assert sfss._kie_fields_for_page(str(canonical), 0) == []
        # Infer xong (kể cả kết quả rỗng) → không còn MISS.
        assert (
            sfss._kie_cached_fields(str(canonical), 0)
            is not sfss._KIE_CACHE_MISS
        )
    finally:
        sfss._KIE_FIELDS_CACHE.clear()


def test_kie_org_fallback_name_ignores_junk_fields(tmp_path, monkeypatch) -> None:
    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        lambda path, page: [
            None,
            "junk",
            42,
            _field("ISSUE_ORG_NAME", "SỞ Y TẾ", page_index=page),
        ],
    )
    assert _kie_org_fallback_name(str(canonical), 0) == "SỞ Y TẾ*"


def test_kie_issue_year_ignores_junk_fields(tmp_path, monkeypatch) -> None:
    from scanindex.ui.screens.secret_file_scan_screen import _kie_issue_year

    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        lambda path, page: [
            None,
            "junk",
            _field("PLACE_DATE", "ngày 10 tháng 10 năm 2004", page_index=page),
        ],
    )
    assert _kie_issue_year(str(canonical), 0) == 2004


def test_safe_page_index_junk() -> None:
    from scanindex.ui.screens.secret_file_scan_screen import _safe_page_index

    assert _safe_page_index(1) == 0
    assert _safe_page_index(3) == 2
    assert _safe_page_index("3") == 2
    assert _safe_page_index(None) == 0
    assert _safe_page_index("rác") == 0
    assert _safe_page_index(-5) == 0


def test_apply_kie_org_fallback_survives_junk_page_number(
    tmp_path, monkeypatch
) -> None:
    canonical = tmp_path / "ocr.json.zst"
    canonical.write_text("{}", encoding="utf-8")
    monkeypatch.setattr(
        "scanindex.ui.screens.secret_file_scan_screen._kie_fields_for_page",
        lambda path, page: [_field("ISSUE_ORG_NAME", "SỞ Y TẾ", page_index=page)],
    )
    m = _match(source_path=_UUID_P9)
    m.page_number = "rác"  # type: ignore[assignment]
    _apply_kie_org_fallback([m], str(canonical))  # không raise
    assert m.kie_org_name == "SỞ Y TẾ*"  # fallback về trang 0


# ---------------------------------------------------------------------------
# Ngày cập nhật file (mtime) — hiển thị GUI + cột Excel
# ---------------------------------------------------------------------------

def test_file_mtime_helpers() -> None:
    # File còn trên đĩa → mtime live.
    live_path = __file__
    m = _match(source_path=live_path)
    assert abs(_file_mtime_for_display(m) - os.stat(live_path).st_mtime) < 1
    # File đã xóa → dùng mtime chụp lúc quét.
    gone = _match(source_path=r"D:\kho\khong-cong.pdf", file_mtime=1758000000.0)
    assert _file_mtime_for_display(gone) == 1758000000.0
    # Không có gì → 0, hiển thị trống.
    nothing = _match(source_path=r"D:\kho\khong-cong.pdf")
    assert _file_mtime_for_display(nothing) == 0.0
    assert _format_file_mtime(0.0) == ""
    assert _format_file_mtime(1758000000.0) != ""


def test_export_writes_file_update_date_cell(tmp_path) -> None:
    """File thật → ô datetime (sắp xếp được trong Excel) + định dạng ngày."""
    from datetime import datetime

    import openpyxl

    real_pdf = tmp_path / "congvan_01.pdf"
    real_pdf.write_bytes(b"%PDF-fake")
    expected = os.stat(str(real_pdf)).st_mtime
    m = _match(source_path=str(real_pdf))

    dest = str(tmp_path / "danh_sach.xlsx")
    export_matches_to_excel([m], dest)

    ws = openpyxl.load_workbook(dest).active
    cell = ws.cell(row=2, column=8)
    assert isinstance(cell.value, datetime)
    assert abs(cell.value.timestamp() - expected) < 1
    assert cell.number_format == "DD/MM/YYYY HH:MM"


def test_load_excel_round_trips_file_update_date(tmp_path) -> None:
    """Nạp lại file xuất ra → file_mtime dựng đúng từ ô ngày (datetime)."""
    real_pdf = tmp_path / "congvan_01.pdf"
    real_pdf.write_bytes(b"%PDF-fake")
    expected = os.stat(str(real_pdf)).st_mtime
    m = _match(source_path=str(real_pdf))

    dest = str(tmp_path / "danh_sach.xlsx")
    export_matches_to_excel([m], dest)

    loaded = load_secret_matches_from_excel(dest)
    assert abs(loaded[0].file_mtime - expected) < 1
    # File còn trên đĩa → hiển thị vẫn lấy live (khớp mtime hiện tại).
    assert abs(
        _file_mtime_for_display(loaded[0]) - os.stat(str(real_pdf)).st_mtime
    ) < 1


def test_load_excel_parses_text_update_date(tmp_path) -> None:
    """Excel cũ/tay ghi ngày dạng chuỗi vẫn parse được về file_mtime."""
    import openpyxl

    src = str(tmp_path / "cu.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    headers = list(EXCEL_HEADERS)
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    ws.append([
        1, "MẬT", "a.pdf", "", "", "a.pdf", r"D:\kho\a.pdf",
        "05/09/2026 14:30", 1, "Trang đầu", "", "",
    ])
    wb.save(src)

    loaded = load_secret_matches_from_excel(src)
    from datetime import datetime as dt

    assert abs(
        loaded[0].file_mtime
        - dt.strptime("05/09/2026 14:30", "%d/%m/%Y %H:%M").timestamp()
    ) < 1
